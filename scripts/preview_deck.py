"""Render the actual .pptx shapes to PNGs (approximate) to verify layout."""
import os
from pptx import Presentation
from pptx.util import Emu
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyArrow
import matplotlib.image as mpimg

BASE=os.path.join(os.path.dirname(__file__),"..")
PPTX=os.path.join(BASE,"reports","Retail_Performance_Review.pptx")
OUT=os.path.join(BASE,"reports","preview"); os.makedirs(OUT,exist_ok=True)
E=914400.0
prs=Presentation(PPTX)
SW=prs.slide_width/E; SH=prs.slide_height/E

def rgb(c):
    try: return "#%02x%02x%02x"%(c[0],c[1],c[2])
    except Exception: return None

for idx,slide in enumerate(prs.slides,1):
    fig,ax=plt.subplots(figsize=(SW,SH),dpi=110)
    ax.set_xlim(0,SW); ax.set_ylim(0,SH); ax.invert_yaxis(); ax.axis("off")
    for sh in slide.shapes:
        try: l,t,w,h=sh.left/E,sh.top/E,sh.width/E,sh.height/E
        except Exception: continue
        if sh.shape_type==13:  # picture
            try:
                img=mpimg.imread(sh.image.blob and __import__("io").BytesIO(sh.image.blob))
                ax.imshow(img,extent=[l,l+w,t+h,t],aspect="auto",zorder=5)
            except Exception as e:
                ax.add_patch(Rectangle((l,t),w,h,facecolor="#dddddd",edgecolor="#999",zorder=5))
            continue
        # shape fill/line
        fill=None; edge=None
        try:
            if sh.fill.type is not None and sh.fill.type==1: fill=rgb(sh.fill.fore_color.rgb)
        except Exception: pass
        try:
            if sh.line.width and sh.line.color and sh.line.color.rgb is not None: edge=rgb(sh.line.color.rgb)
        except Exception: pass
        if fill or edge:
            ax.add_patch(Rectangle((l,t),w,h,facecolor=fill or "none",edgecolor=edge or "none",linewidth=0.8,zorder=2))
        # arrows
        if sh.shape_type is not None and "ARROW" in str(sh.shape_type):
            ax.add_patch(Rectangle((l,t),w,h,facecolor=fill or "#888",zorder=3))
        # text
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf=sh.text_frame
            n=len(tf.paragraphs);
            for pi,p in enumerate(tf.paragraphs):
                if not p.runs: continue
                r=p.runs[0]
                sz=(r.font.size.pt if r.font.size else 12)
                col=rgb(r.font.color.rgb) if (r.font.color and r.font.color.type is not None) else "#0b0b0b"
                try: col=col or "#0b0b0b"
                except: col="#0b0b0b"
                al=p.alignment
                ha={1:"left",2:"center",3:"right"}.get(int(al) if al is not None else 1,"left")
                xx = l if ha=="left" else (l+w/2 if ha=="center" else l+w)
                txt="  ".join(rr.text for rr in p.runs)
                yy=t+0.03+ (pi+0.5)*(h/max(n,1)) if n>1 else t+h/2
                ax.text(xx,yy,txt,fontsize=sz*0.62,color=col,ha=ha,va="center",
                        fontweight="bold" if r.font.bold else "normal",zorder=6,wrap=True)
    ax.add_patch(Rectangle((0,0),SW,SH,fill=False,edgecolor="#cccccc",lw=1))
    fig.subplots_adjust(0,0,1,1)
    fig.savefig(os.path.join(OUT,f"slide{idx:02d}.png"),dpi=110)
    plt.close(fig)
    print("rendered slide",idx)
print("done")
