"""Build the CEO / pipeline deck from the dbt marts (run after `uv run dbt build`).

uv run python scripts/build_deck.py  ->  deliverables/ceo_sales_review.pptx
"""
from pathlib import Path

import duckdb
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
con = duckdb.connect(str(ROOT / "retail.duckdb"), read_only=True)


def q(sql):
    return con.sql(sql).fetchall()


NAVY = RGBColor(0x1F, 0x2A, 0x44)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x1E, 0x84, 0x49)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xEE, 0xF1, 0xF6)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def m(x):
    return f"${x / 1e6:,.2f}M"


def k(x):
    return f"${x / 1e3:,.0f}k"


def pct(x):
    return f"{x * 100:+.1f}%"


def slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text, p.font.size, p.font.bold, p.font.color.rgb = title, Pt(26), True, NAVY
    tb.text_frame.word_wrap = True
    if subtitle:
        p2 = tb.text_frame.add_paragraph()
        p2.text, p2.font.size, p2.font.color.rgb = subtitle, Pt(14), GREY
    return s


def bullets(s, items, x=0.5, y=1.5, w=12.3, h=5.5, size=16):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        level = 1 if it.startswith("- ") else 0
        text = it[2:] if level else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = ("• " if level == 0 else "– ") + text, level
        p.font.size = Pt(size - 2 * level)
        p.space_after = Pt(6)
    return tb


def table(s, rows, x, y, w, col_w=None, size=11, h=0.3):
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h * len(rows)))
    t = shp.table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(size)
                p.font.bold = r == 0
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    return t


def note(s, text):
    s.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- data
kpi = {r[0]: r for r in q("""
    select transaction_year, net_sales, gross_margin, gross_margin_pct, gross_margin_standard, store_opex,
           store_contribution, store_contribution_pct, avg_selling_price, net_sales_yoy_abs, net_sales_yoy_pct,
           gross_margin_yoy_abs, gross_margin_yoy_pct, store_contribution_yoy_abs, store_contribution_yoy_pct,
           avg_selling_price_yoy_pct, units
    from main_marts.mart_kpi_yearly order by 1""")}
k22, k21 = kpi[2022], kpi[2021]
furn = {r[0]: r for r in q("""
    select transaction_year, sum(net_sales), sum(gross_margin), sum(gross_margin_standard), sum(cogs_landed)
    from main_marts.mart_product_yearly where product_category = 'Furniture' group by 1""")}
stores22 = q("""select count(*), sum((store_contribution < 0)::int), sum(case when store_contribution < 0 then store_contribution end)
                from main_marts.mart_store_yearly where transaction_year = 2022 and store_id <> -1""")[0]

# ---------------------------------------------------------------- 1 title
s = slide("2022 Sales Review: growing sales, not growing profit",
          "CEO review + dbt transformation pipeline  |  Data: consolidated_transactions extract, Jan-2019 to Dec-2022")
bullets(s, [
    f"Net Sales {m(k22[1])} in 2022, {pct(k22[10])} YoY — but Gross Margin only {pct(k22[12])}",
    "Two issues the board should see: Furniture is now sold at ~cost, and 63% of stores lose money after direct opex",
    "All figures from dbt marts (fct_sales -> mart_*), landed cost basis, store opex treated as annual",
], y=2.2, size=20)

# ---------------------------------------------------------------- 2 exec summary
s = slide("Executive summary — recommendations first")
f22, f21 = furn[2022], furn[2021]
bullets(s, [
    f"1. Fix Furniture pricing / sourcing now. Landed cost per unit rose ~36% in Jan-2022, prices did not move. "
    f"Furniture GM fell {k(f21[2])} -> {k(f22[2])} ({k(f22[2] - f21[2])}); GM% 32% -> 7%.",
    "- Restoring 2021 margin needs ~+36% price, or a supplier/freight renegotiation; otherwise review range depth.",
    f"2. Report margin on landed cost, not standard cost. On standard cost 2022 GM looks {m(k22[4])} (+8%); "
    f"actual is {m(k22[2])} (+0.6%). The {k(k22[4] - k22[2])} gap is invisible in the budget view.",
    f"3. Store portfolio review. {stores22[1]} of {stores22[0]} stores lose money after direct opex ({k(stores22[2])}). "
    "The West (CA) generates >100% of store contribution; Northeast turned negative in 2022.",
    "- Pause openings outside proven markets: new stores lost $127k in their first year (2022).",
    f"4. Growth is volume-led: +{k(1156073)} volume vs +{k(62724)} price. ASP growth is decelerating (+5.0%, +2.5%, +1.0%).",
    "5. Investigate data/process issues: placeholder 'Unknown' store/product rows, manual-only returns, duplicate ingests.",
], size=16)
note(s, "Lead with the two gems: Furniture landed-cost shock and store profitability. Everything else supports.")

# ---------------------------------------------------------------- 3 pipeline DAG
s = slide("Pipeline: staging -> intermediate -> marts (dbt + DuckDB)",
          "uv run dbt build: 12 models, 23 tests, all passing")
boxes = [
    ("raw\nconsolidated_transactions\n509,480 rows, 24 cols", 0.4, LIGHT),
    ("stg_consolidated_transactions\nsnake_case, types, per-unit costs,\nDQ flags (no rows dropped)", 3.0, LIGHT),
]
inter = ["int_transactions_deduped\n1 row / transaction_id", "int_product_attributes\n1 row / product",
         "int_store_attributes\n1 row / store", "int_transactions_keyed\nkey repair, unknown members"]
marts_core = ["dim_product", "dim_store", "fct_sales\ngrain: transaction line"]
marts_rep = ["mart_kpi_yearly", "mart_product_yearly", "mart_store_yearly", "mart_sales_monthly"]


def box(s, text, x, y, w=2.3, h=0.9, fill=LIGHT):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid()
    b.fill.fore_color.rgb = fill
    b.line.color.rgb = NAVY
    tf = b.text_frame
    tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        p.font.size, p.font.color.rgb = Pt(10), NAVY
    return b


box(s, boxes[0][0], 0.3, 3.2, w=2.2, h=1.1)
box(s, boxes[1][0], 2.8, 3.1, w=2.5, h=1.3)
for i, t in enumerate(inter):
    box(s, t, 5.6, 1.6 + i * 1.3, w=2.4, h=1.0)
for i, t in enumerate(marts_core):
    box(s, t, 8.3, 2.1 + i * 1.3, w=2.0, h=1.0, fill=RGBColor(0xD6, 0xE4, 0xF0))
for i, t in enumerate(marts_rep):
    box(s, t, 10.6, 1.6 + i * 1.3, w=2.4, h=1.0, fill=RGBColor(0xD5, 0xF0, 0xDD))
bullets(s, ["staging: views (cheap, always fresh)  |  intermediate: views  |  marts: tables (what the deck and BI read)"],
        x=0.3, y=6.75, w=12.5, h=0.5, size=12)
for x, label in [(5.6, "intermediate"), (8.3, "marts: star schema"), (10.6, "marts: reporting")]:
    bullets(s, [label], x=x, y=1.15, w=2.5, h=0.4, size=12)

# ---------------------------------------------------------------- 4 normalisation + mart architecture
s = slide("How the extract was normalised & why the mart layer looks like this")
bullets(s, [
    "Grain decision: one fct_sales row per transaction_id (after dedup) = one sales line (month x store x product, qty, price)",
    "- Measures stored additive: net_sales, cogs_landed, cogs_standard, gross_margin(_standard) -> safe to SUM at any level",
    "Dimensions pulled out of the repeated attributes: dim_product (17 + unknown), dim_store (287 + unknown)",
    "- Conformed 'unknown' member (-1) instead of dropping orphans -> relationships tests pass, totals reconcile to source",
    "Two-tier mart layer:",
    "- Star schema (dim_product, dim_store, fct_sales) = reusable source of truth for any BI / ad-hoc question",
    "- Reporting marts at CEO grains (year, product-year, store-year, month x category x region) with YoY abs/%, "
    "price metrics and price/volume bridge pre-computed — the deck reads only these",
    "Why: metric logic (YoY, opex allocation, PVM) lives once in SQL, tested; slides cannot drift from the model",
    "Rejected: one wide 'OBT' mart (reintroduces the attribute conflicts), and putting YoY in Python (untestable, duplicated)",
], size=15)

# ---------------------------------------------------------------- 5 DQ decisions
s = slide("Data-quality issues found and how each was handled", "Nothing silently dropped: every rule is flagged in fct_sales")
table(s, [
    ["Issue", "Evidence", "Decision", "Alternative rejected"],
    ["Duplicate transaction_id", "9,815 ids x2 rows; same date/store/product/qty/price, differ in lineage & cost",
     "Keep latest ingested_at, tie-break POS>ERP>Manual", "Sum both (double counts $0.90M sales)"],
    ["Cost per pack vs per unit", "Dining pack_size=2: unit_cost/price = 117% (sold below cost?)",
     "unit/standard cost ÷ pack_size; landed already per unit", "Multiply by pack_size (makes it worse)"],
    ["3 competing cost fields", "Equal until 2022; Furniture landed +36% vs standard",
     "Landed = margin basis; standard kept as sensitivity", "Standard (Finance) hides the 2022 cost shock"],
    ["Missing keys", "1,310 product ids, 1,308 store ids null; names present",
     "Recover by name (store: name+state if 1:1)", "Drop rows"],
    ["Placeholder 99901", "100 Manual rows Jun-2022, 'Unknown' store+product; 85 ingested before txn date",
     "Map to unknown member -1, flag", "Treat as a real store/product"],
    ["Conflicting store attrs", "Stores 211/212 MD split NE/South; store 50 '(restated)' opex 10x",
     "MD -> South (Census); opex = 4-yr value", "Latest row wins (10x opex error)"],
    ["Product name conflict", "Product 19 'Cosmetics' vs catalog 'Vanities' every year",
     "Keep transaction name, catalog alias retained", "Split into 2 products"],
    ["Negative quantities", "4,907 lines, Manual source only", "Keep as returns (net sales net of returns)", "Exclude"],
], 0.3, 1.35, 12.7, col_w=[2.2, 4.3, 3.4, 2.8], size=10, h=0.5)

# ---------------------------------------------------------------- 6 definitions & sensitivity
s = slide("Definitions and sensitivity to choices")
bullets(s, [
    "Net Sales = quantity x unit_price (returns negative). Gross Margin = Net Sales - quantity x landed cost per unit.",
    "Profitability = Store contribution = Gross Margin - direct store opex (annual, per active store). No central costs in data.",
], size=15, h=1.3)
table(s, [
    ["2022", "Landed cost (chosen)", "Standard cost", "Difference"],
    ["Gross Margin", m(k22[2]), m(k22[4]), k(k22[4] - k22[2])],
    ["GM %", f"{k22[3] * 100:.1f}%", f"{k22[4] / k22[1] * 100:.1f}%", ""],
    ["GM YoY", pct(k22[12]), pct(k22[4] / kpi[2021][4] - 1), ""],
], 0.5, 2.9, 7.5, size=13, h=0.4)
bullets(s, [
    "Opex annual vs monthly: monthly x12 = $40M opex vs $12.8M sales -> every store loses money; not credible -> annual.",
    "Opex is a full-year charge even for stores opened mid-year -> first-year losses slightly overstated.",
    "Dedup rule choice does not change sales or landed-cost GM (both identical within every duplicate pair).",
], y=4.9, size=14, h=2)

# ---------------------------------------------------------------- 7 growth
s = slide("Q1: Sales are growing ~11% a year — profit is not",
          f"2022: Net Sales {pct(k22[10])}, Gross Margin {pct(k22[12])}, Store contribution {pct(k22[14])}")
cd = CategoryChartData()
cd.categories = [str(y) for y in kpi]
cd.add_series("Net Sales ($M)", [round(kpi[y][1] / 1e6, 2) for y in kpi])
cd.add_series("Gross Margin ($M)", [round(kpi[y][2] / 1e6, 2) for y in kpi])
cd.add_series("Store contribution ($M)", [round(kpi[y][6] / 1e6, 2) for y in kpi])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.4), Inches(1.5), Inches(6.6), Inches(5.3), cd).chart
ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
ch.plots[0].has_data_labels = True
ch.plots[0].data_labels.font.size = Pt(9)
rows = [["", "2020", "2021", "2022"]]
for label, i, fmt in [("Net Sales", 1, m), ("  YoY abs", 9, k), ("  YoY %", 10, pct),
                      ("Gross Margin", 2, m), ("  YoY abs", 11, k), ("  YoY %", 12, pct),
                      ("GM %", 3, lambda v: f"{v * 100:.1f}%"),
                      ("Store contribution", 6, m), ("  YoY abs", 13, k), ("  YoY %", 14, pct),
                      ("Avg selling price", 8, lambda v: f"${v:.2f}"), ("  YoY %", 15, pct)]:
    rows.append([label] + [fmt(kpi[y][i]) for y in (2020, 2021, 2022)])
table(s, rows, 7.3, 1.5, 5.7, size=11, h=0.33)

# ---------------------------------------------------------------- 8 GEM 1 furniture
s = slide("Gem #1: Furniture margin collapsed in Jan-2022 — prices were never adjusted",
          "Landed cost per unit +36% overnight; price -0.3%. Standard-cost reporting does not show it.")
mon = q("""select strftime(transaction_month, '%Y-%m'), sum(gross_margin) / sum(net_sales), sum(gross_margin_standard) / sum(net_sales)
           from main_marts.mart_sales_monthly where product_category = 'Furniture' and transaction_month >= '2021-01-01'
           group by 1 order by 1""")
cd = CategoryChartData()
cd.categories = [r[0] for r in mon]
cd.add_series("Furniture GM% — landed cost (actual)", [round(r[1] * 100, 1) for r in mon])
cd.add_series("Furniture GM% — standard cost (budget view)", [round(r[2] * 100, 1) for r in mon])
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.4), Inches(1.5), Inches(7.6), Inches(5.3), cd).chart
ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
fp = q("""select product_name, gross_margin_pct, gross_margin, gross_margin_py, price_yoy_pct, landed_cost_yoy_pct
          from main_marts.mart_product_yearly where product_category = 'Furniture' and transaction_year = 2022 order by net_sales desc""")
table(s, [["2022", "GM%", "GM", "GM 2021", "Price YoY", "Cost YoY"]] +
      [[r[0], f"{r[1] * 100:.1f}%", k(r[2]), k(r[3]), pct(r[4]), pct(r[5])] for r in fp],
      8.2, 1.5, 4.9, col_w=[1.5, 0.6, 0.7, 0.75, 0.7, 0.65], size=10, h=0.4)
bullets(s, [
    f"Furniture: {k(f22[1])} sales (+10%) but GM {k(f21[2])} -> {k(f22[2])}",
    "Explains the whole gap between sales growth (+10.7%) and GM growth (+0.6%)",
    "Action: reprice (~+36% to restore), renegotiate freight/supplier, or cut range; flag to Finance",
], x=8.2, y=3.4, w=4.9, h=3.5, size=13)

# ---------------------------------------------------------------- 9 drivers & price
s = slide("Q3: 2022 growth drivers — volume, not price", "Price/volume bridge from mart_product_yearly")
cat = q("""select product_category, sum(price_effect), sum(volume_effect), sum(net_sales_yoy_abs),
                  sum(net_sales) / sum(units), sum(gross_margin) / sum(units),
                  sum(gross_margin_py) / sum(units_py)
           from main_marts.mart_product_yearly where transaction_year = 2022 and product_id <> -1
           group by 1 order by 4 desc""")
cd = CategoryChartData()
cd.categories = [r[0] for r in cat]
cd.add_series("Price effect ($k)", [round(r[1] / 1e3) for r in cat])
cd.add_series("Volume effect ($k)", [round(r[2] / 1e3) for r in cat])
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.4), Inches(1.5), Inches(6.6), Inches(5.3), cd).chart
ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
ch.plots[0].has_data_labels = True
ch.plots[0].data_labels.font.size = Pt(9)
bullets(s, [
    "Total 2022 Net Sales +$1.22M = +$1.16M volume (+9.6% units) + $0.06M price",
    "Existing stores delivered +$1.12M; 14 new stores only $0.09M",
    "Price derivatives:",
    "- ASP growth decelerating: +5.0% (2020), +2.5%, +1.0% (2022)",
    "- Electronics cut prices 5-7% in line with 7% cost decline -> GM% held, volume +12%",
    "- Household/Kitchen/Dining/Sports: price +3-4% vs landed cost +2.9% -> healthy price-cost spread",
    "- Furniture: price -0.3% vs cost +36% -> spread turned to ~zero",
    "- GM per unit 2022 vs 2021: Furniture fell most; track 'price-cost spread' monthly per category",
], x=7.2, y=1.5, w=5.9, h=5.5, size=13)

# ---------------------------------------------------------------- 10 GEM 2 stores/regions
s = slide(f"Gem #2: {stores22[1]} of {stores22[0]} stores lose money after direct opex",
          "The West (mainly California) carries the business; Northeast turned negative in 2022")
reg = q("""select region_name, sum(net_sales), sum(store_contribution), count(*), sum((not is_profitable)::int)
           from main_marts.mart_store_yearly where transaction_year = 2022 and store_id <> -1 group by 1 order by 2 desc""")
cd = CategoryChartData()
cd.categories = [r[0] for r in reg]
cd.add_series("Net Sales 2022 ($k)", [round(r[1] / 1e3) for r in reg])
cd.add_series("Store contribution 2022 ($k)", [round(r[2] / 1e3) for r in reg])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.4), Inches(1.5), Inches(6.2), Inches(4.0), cd).chart
ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
ch.plots[0].has_data_labels = True
ch.plots[0].data_labels.font.size = Pt(9)
table(s, [["Region", "Stores", "Loss-making"]] + [[r[0], r[3], r[4]] for r in reg], 0.6, 5.6, 5.8, size=11, h=0.3)
top = q("""select store_name || ', ' || state_code, net_sales, store_contribution, net_sales_yoy_pct from main_marts.mart_store_yearly
           where transaction_year = 2022 and store_id <> -1 order by store_contribution desc limit 5""")
table(s, [["Top stores 2022", "Sales", "Contrib.", "YoY"]] + [[r[0], k(r[1]), k(r[2]), pct(r[3])] for r in top],
      6.9, 1.5, 6.1, col_w=[2.6, 1.2, 1.2, 1.1], size=11, h=0.33)
bullets(s, [
    "Top products 2022: Computers $2.0M, Dinnerware $1.8M, Cookware $1.7M (by GM: Dinnerware, Candles, Cookware)",
    "Los Angeles alone = 8% of sales; top 4 CA stores = $1.2M contribution",
    "NYC sales -10.5%, Northeast sales -9% in 2022 — investigate (competition? store issues?)",
    "Long tail: many small Township/Borough stores with opex > gross margin",
], x=6.9, y=3.7, w=6.1, h=3.2, size=13)

# ---------------------------------------------------------------- 11 profit levers + issues
s = slide("Q4 & Q5: How to grow profit, and what needs investigation")
bullets(s, [
    "Profit levers (size of prize, 2022 base)",
    "- Furniture price/cost fix: up to ~$0.49M GM (back to 2021 GM%)",
    "- Store portfolio: 182 loss-making stores cost $1.0M; close/relocate worst decile, renegotiate leases, or test smaller formats",
    "- Pricing discipline: pass through cost inflation category by category (price-cost spread KPI)",
    "- Focus growth capital on West/CA format that works; gate new openings on a payback model",
    "Issues to investigate",
    "- 'Unknown' store/product 99901: 100 Manual lines, Jun-2022, $12.9k, ingested before the sale date",
    "- Returns only ever recorded via Manual source — POS/ERP returns missing? Returns value +21% in 2022",
    "- 9,815 duplicate ingests across systems; cost values differ between copies",
    "- Store master data: region conflicts (MD), 10x opex restatement (store 50), duplicate store names (Wayne, IN)",
], size=15)

# ---------------------------------------------------------------- 12 assumptions / next steps
s = slide("Assumptions, trade-offs, and what I would do with more time")
bullets(s, [
    "Assumptions: monthly grain is the economic date; landed cost = true cost; opex annual & fully charged to active stores",
    "Trade-offs: mode() for conflicting attributes (simple, explainable) instead of SCD2 history; views for staging/intermediate",
    "Not done (time): incremental fct_sales, SCD2 dim_store / dim_product with effective dates, unit tests on dedup rule,",
    "- dbt source freshness & DQ dashboards, mix effect split inside volume, store-level like-for-like with partial years",
    "Next: confirm Furniture landed-cost change with AP/Ops; agree one margin definition with Finance; store P&L review",
], size=16)

out = ROOT / "deliverables" / "ceo_sales_review.pptx"
out.parent.mkdir(exist_ok=True)
prs.save(out)
print(out)
